"""Excel rendering of a report: one sheet per section with raw numbers (bytes, bit/s, %),
for teams who build their own charts. An "Info" sheet describes the period and the scope."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="F58220")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FDF1E6")


def _naive(v: Any, tz: ZoneInfo) -> Any:
    """Excel has no time zones: datetimes in the report time zone, without tzinfo."""
    if isinstance(v, datetime):
        d = v if v.tzinfo else v.replace(tzinfo=timezone.utc)
        return d.astimezone(tz).replace(tzinfo=None)
    if isinstance(v, str) and len(v) >= 19 and v[4] == "-" and v[10] == "T":
        try:
            return _naive(datetime.fromisoformat(v), tz)
        except ValueError:
            return v
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return v


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, str, str | None]], rows: list[dict[str, Any]],
           tz: ZoneInfo, warn_key: str | None = None) -> None:
    """columns: (header, key, number format)."""
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_naive(r.get(k), tz) for _, k, _ in columns])
        if warn_key and r.get(warn_key):
            for cell in ws[ws.max_row]:
                cell.fill = WARN_FILL
    for i, (h, _, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        width = max([len(h)] + [len(str(c.value)) for c in ws[letter][1:200] if c.value is not None]) + 2
        ws.column_dimensions[letter].width = min(max(width, 9), 48)
        if fmt:
            for c in ws[letter][1:]:
                c.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def render_xlsx(r: dict[str, Any]) -> bytes:
    tz = ZoneInfo(r["tz"])
    wb = Workbook()
    info = wb.active
    info.title = "Info"
    rows = [("Report", r["title"]), ("Period", r["label"]), ("From", _naive(r["start"], tz)), ("To", _naive(r["end"], tz)),
            ("Compared with", f'{_naive(r["previous_start"], tz):%Y-%m-%d} → {_naive(r["previous_end"], tz):%Y-%m-%d}'),
            ("Time zone", r["tz"]), ("Scope", ", ".join(r["scope"])), ("Generated", _naive(r["generated_at"], tz)),
            ("Units", "bytes, bit/s, packets; percent = share or utilization"),
            ("Note", "Traffic volumes are estimates (sFlow sampling); utilization comes from interface counters (exact).")]
    for k, v in rows:
        info.append([k, v])
        info.cell(info.max_row, 1).font = Font(bold=True)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 16, 80
    for c in info["B"]:
        if isinstance(c.value, datetime):
            c.number_format = "yyyy-mm-dd hh:mm"
            c.alignment = Alignment(horizontal="left")

    B, BPS, PCT, INT, DT = "#,##0", "#,##0", "0.00", "#,##0", "yyyy-mm-dd hh:mm"
    if "summary" in r:
        s = r["summary"]
        _sheet(wb, "Summary", [("Indicator", "k", None), ("This period", "v", "#,##0.##"), ("Previous period", "p", "#,##0.##"),
                               ("Change %", "c", "0.0")], [
            {"k": "Traffic volume (bytes, estimated)", "v": s["bytes"], "p": s["prev_bytes"] if s["has_previous"] else None, "c": s["bytes_change"]},
            {"k": "Average throughput (bit/s)", "v": s["bps"], "p": s["prev_bps"] if s["has_previous"] else None, "c": s["bps_change"]},
            {"k": "Switches sending sFlow", "v": s["exporters"]}, {"k": "Source hosts", "v": s["sources"]},
            {"k": "Destination hosts", "v": s["destinations"]}, {"k": "Conversations", "v": s["conversations"]},
            {"k": "sFlow samples", "v": s["samples"]}, {"k": "Alerts", "v": s["alerts"]},
            {"k": "Broadcast storms", "v": s["storm_events"]}, {"k": "Ports with broadcast storms", "v": s["storm_ports"]}], tz)
    if "ports" in r:
        cols = [("Group", "group", None), ("Switch", "exporter_name", None), ("Port", "label", None), ("Interface id", "interface_id", None),
                ("Speed (bit/s)", "speed_bps", INT), ("Avg in (bit/s)", "in_avg_bps", BPS), ("Avg out (bit/s)", "out_avg_bps", BPS),
                ("Avg in %", "in_avg_pct", PCT), ("Avg out %", "out_avg_pct", PCT), ("P95 in %", "in_p95_pct", PCT),
                ("P95 out %", "out_p95_pct", PCT), ("P95 %", "p95_pct", PCT), ("P95 % previous", "prev_p95_pct", PCT),
                ("Peak in %", "in_max_pct", PCT), ("Peak out %", "out_max_pct", PCT), ("Discards in", "in_discards", INT),
                ("Discards out", "out_discards", INT), ("Errors in", "in_errors", INT), ("Errors out", "out_errors", INT),
                ("Above threshold", "warn", None)]
        if r["ports"]["monitored"]:
            _sheet(wb, "Monitored links", cols, r["ports"]["monitored"], tz, "warn")
        if r["ports"].get("volumes"):
            vcols = [("Group", "group", None), ("Switch", "exporter_name", None), ("Port", "label", None),
                     ("Interface id", "interface_id", None), ("Bytes", "bytes", B), ("Previous period bytes", "prev_bytes", B),
                     ("Change %", "prev_change", "0.0")]
            if r["ports"]["volumes_year"]:
                vcols += [("Same period last year bytes", "year_bytes", B), ("Change vs last year %", "year_change", "0.0")]
            _sheet(wb, "Monitored volumes", vcols, r["ports"]["volumes"], tz)
        _sheet(wb, "Busiest ports", cols[1:], r["ports"]["busiest"], tz, "warn")
    if "applications" in r:
        _sheet(wb, "Applications", [("Application", "service", None), ("Bytes", "bytes", B), ("Packets", "packets", INT),
                                    ("Samples", "samples", INT), ("Share %", "percent", PCT)], r["applications"], tz)
        _sheet(wb, "Protocols", [("Protocol", "protocol", None), ("Bytes", "bytes", B), ("Packets", "packets", INT),
                                 ("Samples", "samples", INT), ("Share %", "percent", PCT)], r["protocols"], tz)
    if "sources" in r:
        host_cols = [("IP", "ip", None), ("Name", "name", None), ("Vendor", "vendor", None), ("MAC", "mac", None),
                     ("Bytes", "bytes", B), ("Packets", "packets", INT), ("Share %", "percent", PCT)]
        _sheet(wb, "Top sources", host_cols, r["sources"], tz)
        _sheet(wb, "Top destinations", host_cols, r["destinations"], tz)
    if "conversations" in r:
        _sheet(wb, "Conversations", [("Host A", "host_a", None), ("Name A", "host_a_name", None), ("Host B", "host_b", None),
                                     ("Name B", "host_b_name", None), ("Protocol", "protocol", None), ("Application", "service", None),
                                     ("A → B bytes", "a_to_b_bytes", B), ("B → A bytes", "b_to_a_bytes", B), ("Total bytes", "bytes", B),
                                     ("Share %", "percent", PCT), ("VLANs", "vlans", None), ("First seen", "first_seen", DT),
                                     ("Last seen", "last_seen", DT)], r["conversations"], tz)
    if "trend" in r:
        _sheet(wb, "Trend", [("Bucket", "label", None), ("Bytes", "bytes", B), ("Average bit/s", "bps", BPS),
                             ("Previous period bytes", "prev_bytes", B)], r["trend"]["points"], tz)
    if "alerts" in r:
        _sheet(wb, "Alerts", [("Severity", "severity", None), ("Rule", "rule", None), ("On", "subject", None),
                              ("State", "state", None), ("Opened", "opened_at", DT), ("Closed", "closed_at", DT),
                              ("Duration (s)", "duration_s", INT), ("Value", "value", "#,##0.##"), ("Details", "message", None)],
               r["alerts"], tz)
        if r["storms"]:
            _sheet(wb, "Broadcast storms", [("Switch", "exporter_name", None), ("Port", "label", None),
                                            ("Peak bcast in (pps)", "in_bcast_max_pps", "#,##0"),
                                            ("Peak bcast out (pps)", "out_bcast_max_pps", "#,##0"), ("Peak at", "bcast_peak_at", DT),
                                            ("Polls above threshold", "storm_polls", INT)], r["storms"], tz)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
