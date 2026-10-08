"""Reports: definitions, schedule, generation (PDF + Excel) against the test dataset, e-mail, history, roles."""

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest

from app import notify, store
from app.reports import service
from app.reports.data import SECTIONS

from .conftest import EXP_A, exp_id

UI = {"X-Requested-With": "sflow"}
PARIS = ZoneInfo("Europe/Paris")
# test dataset: 2026-01-01 (Thursday) → week of Mon 2025-12-29, reported on Mon 2026-01-05
MONDAY_9H = datetime(2026, 1, 5, 9, 0, tzinfo=PARIS)


@pytest.fixture()
def clean():
    store.execute("DELETE FROM reports")
    store.execute("DELETE FROM report_definitions")
    yield
    store.execute("DELETE FROM reports")
    store.execute("DELETE FROM report_definitions")


def _def(**kw):
    d = {"name": "Weekly network report", "period": "weekly",
         "options": {"monitored_interfaces": [f"{exp_id(EXP_A)}/48"]}, "delivery": {"formats": ["pdf", "xlsx"]}}
    d.update(kw)
    return service.save_definition(d)


def test_validation():
    v = service.validate({"name": " Uplinks ", "period": "monthly",
                          "options": {"sections": ["alerts", "summary"], "busiest_ports": "5", "exporters": [" a ", "a"]},
                          "delivery": {"formats": ["xlsx", "csv"], "recipients": ["noc@example.com"]}})
    assert v["name"] == "Uplinks" and v["options"]["sections"] == ["summary", "alerts"]  # canonical order
    assert v["options"]["busiest_ports"] == 5 and v["options"]["exporters"] == ["a"]
    assert v["delivery"]["formats"] == ["xlsx"]
    for bad in ({"period": "daily"}, {"options": {"sections": ["nope"]}}, {"options": {"sections": []}},
                {"options": {"warn_pct": 150}}, {"delivery": {"formats": []}},
                {"delivery": {"recipients": ["not-an-address"]}}, {"name": ""}):
        with pytest.raises(service.ReportError):
            service.validate({"name": "x", "period": "weekly", **bad})


def test_schedule_weekly_and_monthly(clean):
    d = _def()
    # created now: the last complete week is considered done, the next one is scheduled
    assert d["last_period"] and d["next_run"]
    store.execute("UPDATE report_definitions SET last_period = 'old'")
    assert service.due(MONDAY_9H - timedelta(hours=2)) == []       # Monday 07:00: not yet
    due = service.due(MONDAY_9H)
    assert [x["id"] for x in due] == [d["id"]]
    assert due[0]["_key"] == datetime(2025, 12, 29, tzinfo=PARIS).isoformat()
    m = service.save_definition({"name": "Monthly", "period": "monthly"})
    store.execute("UPDATE report_definitions SET last_period = 'old' WHERE id = ?", (m["id"],))
    keys = {x["name"]: x["_key"] for x in service.due(datetime(2026, 1, 1, 8, 30, tzinfo=PARIS))}
    assert keys["Monthly"] == datetime(2025, 12, 1, tzinfo=PARIS).isoformat()
    store.execute("UPDATE report_definitions SET enabled = 0")
    assert service.due(MONDAY_9H) == []


def test_scheduled_generation_and_email(testdb, clean, monkeypatch):
    sent = []

    def fake_send(subject, body, recipients=None, attachments=()):
        sent.append((subject, body, recipients, [(n, m, len(c)) for n, m, c in attachments]))
        return recipients or ["default@example.com"]

    monkeypatch.setattr(notify, "send_mail", fake_send)
    monkeypatch.setattr(notify, "settings", lambda: {"email": {"enabled": True}, "public_url": "https://sflow.example"})
    d = _def(delivery={"formats": ["pdf", "xlsx"], "email": True, "recipients": ["noc@example.com"]})
    store.execute("UPDATE report_definitions SET last_period = 'old'")
    r = service.run_scheduled(testdb, MONDAY_9H)
    assert len(r["generated"]) == 1 and not r["errors"], r
    assert service.due(MONDAY_9H) == []                       # done once
    rep = service.list_reports()[0]
    assert rep["status"] == "ok" and rep["error"] is None and rep["definition_id"] == d["id"]
    assert rep["trigger"] == "schedule" and rep["period_start"].startswith("2025-12-29")
    assert rep["summary"]["bytes"] > 0
    pdf, name = service.report_file(rep["id"], "pdf")
    assert pdf.read_bytes()[:5] == b"%PDF-" and name == "weekly-network-report-2025-12-29.pdf"
    xlsx, _ = service.report_file(rep["id"], "xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xlsx.read_bytes()))
    assert {"Info", "Summary", "Monitored links", "Busiest ports", "Applications"} <= set(wb.sheetnames)
    assert len(sent) == 1
    subject, body, to, att = sent[0]
    assert "Weekly network report" in subject and to == ["noc@example.com"]
    assert [a[0] for a in att] == ["weekly-network-report-2025-12-29.pdf", "weekly-network-report-2025-12-29.xlsx"]
    assert "Traffic volume" in body and "https://sflow.example/reports" in body
    assert rep["delivery"]["email"]["result"] == "ok"


def test_email_failure_is_recorded(testdb, clean, monkeypatch):
    monkeypatch.setattr(notify, "settings", lambda: {"email": {"enabled": False}, "public_url": ""})
    d = _def(options={"sections": ["summary"]}, delivery={"formats": ["pdf"], "email": True})
    rep = service.generate(testdb, service.queue(service.get_definition(d["id"]), "manual", "admin")["id"])
    assert rep["status"] == "ok" and "not sent" in rep["error"] and rep["files"].keys() == {"pdf"}


def test_custom_period_and_sections(testdb, clean):
    d = _def(options={"sections": ["summary", "trends"]}, delivery={"formats": ["pdf"]})
    s, e = service.custom_bounds(date(2026, 1, 1), date(2026, 1, 1))
    assert (e - s) == timedelta(days=1)
    rep = service.queue(service.get_definition(d["id"]), "manual", "admin", "custom", s, e)
    rep = service.generate(testdb, rep["id"])
    assert rep["status"] == "ok" and rep["period"] == "custom" and rep["label"].startswith("2026-01-01")
    with pytest.raises(service.ReportError):
        service.custom_bounds(date(2026, 1, 2), date(2026, 1, 1))
    with pytest.raises(service.ReportError):
        service.custom_bounds(date(2024, 1, 1), date(2026, 1, 1))


def test_html_numbers_sections():
    from app.reports.pdf import render_html

    r = {"title": "T", "label": "L", "tz": "UTC", "start": datetime(2026, 1, 1, tzinfo=timezone.utc),
         "end": datetime(2026, 1, 2, tzinfo=timezone.utc), "previous_start": datetime(2025, 12, 31, tzinfo=timezone.utc),
         "previous_end": datetime(2026, 1, 1, tzinfo=timezone.utc), "generated_at": datetime(2026, 1, 2, tzinfo=timezone.utc),
         "sections": ["alerts"], "scope": ["all switches"], "alerts": [], "storms": [], "storm_events": []}
    html = render_html(r)
    assert "1. Alerts and incidents" in html and "No alert during this period" in html


def test_purge(clean, tmp_path, monkeypatch):
    monkeypatch.setattr(service, "reports_dir", lambda: tmp_path)
    old = (datetime.now(timezone.utc) - timedelta(days=service.KEEP_DAYS + 1)).isoformat(timespec="seconds")
    rid = store.connection().execute(
        "INSERT INTO reports (title, period, trigger, created_at, status, files) VALUES ('x', 'weekly', 'manual', ?, 'ok', "
        "'{\"pdf\": {\"size\": 1}}')", (old,)).lastrowid
    (tmp_path / f"{rid}.pdf").write_bytes(b"x")
    (tmp_path / "999999.xlsx").write_bytes(b"orphan")
    assert service.purge() == 1
    assert list(tmp_path.iterdir()) == [] and service.list_reports() == []


def test_api_definitions_run_and_download(client, clean):
    defs = client.get("/api/v1/report-definitions").json()
    assert defs["sections"] == list(SECTIONS) and defs["send_hour"] == 8
    r = client.post("/api/v1/report-definitions", headers=UI,
                    json={"name": "API report", "period": "weekly", "options": {"sections": ["summary", "applications"]},
                          "delivery": {"formats": ["xlsx"]}})
    assert r.status_code == 200, r.text
    did = r.json()["id"]
    assert client.post("/api/v1/report-definitions", headers=UI, json={"name": "x", "period": "daily"}).status_code == 422
    r = client.post(f"/api/v1/report-definitions/{did}/run", headers=UI,
                    json={"period": "custom", "start": "2026-01-01", "end": "2026-01-01"})
    assert r.status_code == 202, r.text
    rid = r.json()["id"]
    # TestClient runs the background task before returning
    items = client.get("/api/v1/reports").json()["items"]
    assert items[0]["id"] == rid and items[0]["status"] == "ok" and "delivery" not in items[0]
    f = client.get(f"/api/v1/reports/{rid}/xlsx")
    assert f.status_code == 200 and f.content[:2] == b"PK"
    assert "api-report-2026-01-01.xlsx" in f.headers["content-disposition"]
    assert client.get(f"/api/v1/reports/{rid}/pdf").status_code == 404     # not produced
    assert client.post(f"/api/v1/report-definitions/{did}/run", headers=UI,
                       json={"period": "custom", "start": "2026-01-01"}).status_code == 422
    assert client.delete(f"/api/v1/report-definitions/{did}", headers=UI).status_code == 200
    assert client.get("/api/v1/reports").json()["items"][0]["definition_id"] is None   # history kept
    assert client.delete(f"/api/v1/reports/{rid}", headers=UI).status_code == 200
    assert client.get(f"/api/v1/reports/{rid}/xlsx").status_code == 404
